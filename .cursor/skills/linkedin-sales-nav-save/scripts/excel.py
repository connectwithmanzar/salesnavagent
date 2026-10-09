from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from parse import safe_filename, split_name

COLUMNS = [
    "First Name",
    "Last Name",
    "Full Name",
    "Title",
    "Company",
    "Location",
    "Sales Navigator URL",
    "Date added",
]


def write_files(leads: list[dict], list_name: str, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = []
    for lead in leads:
        first, last = split_name(lead.get("name") or "")
        rows.append(
            {
                "First Name": first,
                "Last Name": last,
                "Full Name": lead.get("name") or "",
                "Title": lead.get("title") or "",
                "Company": lead.get("company") or "",
                "Location": lead.get("location") or "",
                "Sales Navigator URL": lead.get("linkedin_url") or "",
                "Date added": lead.get("extra") or "",
            }
        )

    stem = safe_filename((list_name or "").replace("Lead Lists", "").strip() or "SN_people_list")
    out_xlsx = out_dir / f"{stem}.xlsx"
    out_csv = out_dir / f"{stem}.csv"

    with out_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (list_name or "Leads")[:31]
    navy = PatternFill("solid", fgColor="1B365D")
    white = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body = Font(name="Arial", size=10)
    wrap = Alignment(vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    alt = PatternFill("solid", fgColor="F8FAFC")
    sheet.append(COLUMNS)
    for col in range(1, len(COLUMNS) + 1):
        cell = sheet.cell(1, col)
        cell.fill = navy
        cell.font = white
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append([row[col] for col in COLUMNS])
        row_idx = sheet.max_row
        for col in range(1, len(COLUMNS) + 1):
            cell = sheet.cell(row_idx, col)
            cell.font = body
            cell.alignment = wrap
            cell.border = thin
            if row_idx % 2 == 0:
                cell.fill = alt
        sheet.row_dimensions[row_idx].height = 22
    last = max(sheet.max_row, 1)
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    for col, width in zip("ABCDEFGH", [16, 20, 26, 44, 24, 28, 72, 14]):
        sheet.column_dimensions[col].width = width
    sheet.sheet_view.showGridLines = False
    workbook.properties.title = list_name or "Sales Navigator list"
    workbook.save(out_xlsx)
    print(f"wrote {len(rows)} leads -> {out_xlsx}", flush=True)
    return out_xlsx
