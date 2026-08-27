#!/usr/bin/env python3
"""Five-slide playbook: Sales Navigator list → CSV."""

from pathlib import Path

from PIL import Image as PILImage
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ASSETS = Path("/Users/manzar.i/Desktop/sdr-data-consolidator/presentations/assets")
OUT = Path("/Users/manzar.i/Desktop/sdr-data-consolidator/presentations/Taboola_SN_Export_Playbook.pptx")

BLUE = RGBColor(0x00, 0x56, 0xF0)
NAVY = RGBColor(0x00, 0x28, 0x52)
INK = RGBColor(0x00, 0x1D, 0x37)
MUTED = RGBColor(0x4A, 0x5E, 0x73)
LINE = RGBColor(0xD5, 0xDE, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xF7, 0xF9, 0xFC)
PILL = RGBColor(0xE8, 0xF0, 0xFF)
DEEP = RGBColor(0x00, 0x33, 0x66)

W = Inches(13.333)
H = Inches(7.5)
TOTAL = 5


def _run(p, text, size, bold, color):
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"


def rect(slide, l, t, w, h, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    return sh


def round_box(slide, l, t, w, h, fill, line=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def txt(slide, l, t, w, h, text, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT, anchor="t"):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", anchor)
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    _run(p, text, size, bold, color)
    return tb


def logos(slide, dark=False):
    t_logo = ASSETS / ("taboola-white-lockup.png" if dark else "taboola-blue-lockup.png")
    c_logo = ASSETS / ("connexity-white-lockup.png" if dark else "connexity-dark-lockup.png")
    h = Inches(0.28)
    tw = h * (643 / 160)
    left = Inches(0.6)
    top = Inches(0.32)
    slide.shapes.add_picture(str(t_logo), left, top, height=h)
    div = RGBColor(0x7A, 0x90, 0xA8) if dark else LINE
    rect(slide, left + tw + Inches(0.14), Inches(0.38), Inches(0.015), Inches(0.16), div)
    slide.shapes.add_picture(str(c_logo), left + tw + Inches(0.26), top, height=h)


def footer(slide, n, dark=False):
    c = RGBColor(0x9B, 0xB0, 0xC4) if dark else MUTED
    txt(slide, Inches(0.6), Inches(7.12), Inches(8.8), Inches(0.24),
        "Internal  ·  prepared by Manzar Imam", 11, False, c)
    txt(slide, Inches(10.5), Inches(7.12), Inches(2.2), Inches(0.24),
        f"{n}  /  {TOTAL}", 11, False, c, PP_ALIGN.RIGHT)


def light_bg(slide):
    rect(slide, 0, 0, W, H, SOFT)
    rect(slide, 0, 0, Inches(0.1), H, BLUE)


def dark_bg(slide):
    rect(slide, 0, 0, W, H, NAVY)
    rect(slide, 0, 0, Inches(0.1), H, BLUE)


def chevron(slide, l, t):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, l, t, Inches(0.28), Inches(0.18))
    sh.fill.solid()
    sh.fill.fore_color.rgb = BLUE
    sh.line.fill.background()


def number_circle(slide, l, t, n, fill=BLUE, size=0.46):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, Inches(size), Inches(size))
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
    txt(slide, l, t + Inches(0.05), Inches(size), Inches(size - 0.08),
        str(n), 16, True, WHITE, PP_ALIGN.CENTER)


def pic(slide, name, l, t, width, max_h=None):
    path = ASSETS / name
    with PILImage.open(path) as im:
        ratio = im.height / im.width
    w = width
    h = width * ratio
    if max_h is not None and h > max_h:
        h = max_h
        w = max_h / ratio
    slide.shapes.add_picture(str(path), l, t, width=w, height=h)


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1 Cover
    s = prs.slides.add_slide(blank)
    dark_bg(s)
    logos(s, dark=True)
    txt(s, Inches(0.6), Inches(1.7), Inches(12), Inches(0.35),
        "INTERNAL PLAYBOOK", 12, True, RGBColor(0x7E, 0xC8, 0xFF))
    txt(s, Inches(0.6), Inches(2.15), Inches(12.1), Inches(1.8),
        "How to export a Sales Navigator list to CSV", 40, True, WHITE)
    txt(s, Inches(0.6), Inches(4.15), Inches(12), Inches(0.5),
        "Open the list. Ask Cursor. Get the file.", 20, False, RGBColor(0xC5, 0xD4, 0xE4))
    rect(s, Inches(0.6), Inches(5.05), Inches(1.1), Inches(0.06), BLUE)
    txt(s, Inches(0.6), Inches(5.3), Inches(8), Inches(0.35), "Manzar Imam", 22, True, WHITE)
    txt(s, Inches(0.6), Inches(5.7), Inches(10), Inches(0.35),
        "SDR  ·  Taboola  ·  Connexity", 16, False, RGBColor(0x9B, 0xB0, 0xC4))
    footer(s, 1, dark=True)

    # 2 Problem + solution + impact
    s = prs.slides.add_slide(blank)
    light_bg(s)
    logos(s)
    txt(s, Inches(0.6), Inches(0.82), Inches(12), Inches(0.28), "WHY THIS EXISTS", 12, True, BLUE)
    txt(s, Inches(0.6), Inches(1.12), Inches(12.1), Inches(0.7),
        "Finding leads in Sales Nav is the job. Pulling a bulk list out is the grind.",
        24, True, INK)

    round_box(s, Inches(0.6), Inches(2.0), Inches(4.0), Inches(4.7), WHITE, LINE)
    txt(s, Inches(0.85), Inches(2.18), Inches(3.5), Inches(0.32), "The problem", 13, True, RGBColor(0xC0, 0x45, 0x2A))
    txt(s, Inches(0.85), Inches(2.55), Inches(3.5), Inches(3.9),
        "You search, filter and save 50–130 people into a list. Then you still need them in a sheet or CRM.\n\nThere is often no Download. 25 names per page. Copy-paste into Excel. Titles break. LinkedIn links get lost. Easy to skip a page.\n\nA 129-person list is six pages of clicking — time you are not writing InMails.",
        14, False, MUTED)

    round_box(s, Inches(4.8), Inches(2.0), Inches(4.0), Inches(4.7), WHITE, LINE)
    rect(s, Inches(4.8), Inches(2.0), Inches(0.1), Inches(4.7), BLUE)
    txt(s, Inches(5.15), Inches(2.18), Inches(3.4), Inches(0.32), "How we solve it", 13, True, BLUE)
    txt(s, Inches(5.15), Inches(2.55), Inches(3.4), Inches(3.9),
        "Leave the saved list open in Chrome.\n\nIn Cursor Agent, send one sentence.\n\nCursor reads every page and writes a CSV to Downloads: name, title, company, location, LinkedIn URL.\n\nNo Chrome extension. No copy-paste.",
        14, False, MUTED)

    round_box(s, Inches(9.0), Inches(2.0), Inches(3.75), Inches(4.7), NAVY)
    txt(s, Inches(9.25), Inches(2.18), Inches(3.3), Inches(0.32), "The impact", 13, True, RGBColor(0x7E, 0xC8, 0xFF))
    txt(s, Inches(9.25), Inches(2.6), Inches(3.3), Inches(3.8),
        "Minutes, not a half-hour of clicking.\n\nFull list, not a partial one.\nNicole 2: 129/129\nSIMILAR MERCHENT V2: 49/49\n\nSame file goes into Sheets, Salesforce or a sequence the same day. Repeat on every new list.",
        14, False, WHITE)
    footer(s, 2)

    # 3 What we need + screenshot
    s = prs.slides.add_slide(blank)
    light_bg(s)
    logos(s)
    txt(s, Inches(0.6), Inches(0.8), Inches(12), Inches(0.26), "BEFORE YOU START", 12, True, BLUE)
    txt(s, Inches(0.6), Inches(1.08), Inches(12), Inches(0.45), "You only need four things.", 24, True, INK)
    needs = [
        ("1  Chrome", "Logged in to Sales Navigator"),
        ("2  A saved list", "Open in a Lead lists tab"),
        ("3  Cursor", "Install once from cursor.com"),
        ("4  One Chrome tick", "Shown on the next slide"),
    ]
    for i, (title, body) in enumerate(needs):
        x = Inches(0.6 + i * 3.15)
        round_box(s, x, Inches(1.62), Inches(3.0), Inches(1.15), WHITE, LINE)
        txt(s, x + Inches(0.16), Inches(1.72), Inches(2.7), Inches(0.4), title, 14, True, INK)
        txt(s, x + Inches(0.16), Inches(2.12), Inches(2.7), Inches(0.48), body, 12, False, MUTED)
    round_box(s, Inches(0.6), Inches(2.95), Inches(12.15), Inches(3.9), WHITE, LINE)
    pic(s, "shot_sn_list.png", Inches(0.85), Inches(3.1), Inches(11.65), max_h=Inches(3.55))
    footer(s, 3)

    # 4 Setup + screenshots
    s = prs.slides.add_slide(blank)
    light_bg(s)
    logos(s)
    txt(s, Inches(0.6), Inches(0.8), Inches(12), Inches(0.26), "ONE-TIME SETUP", 12, True, BLUE)
    txt(s, Inches(0.6), Inches(1.08), Inches(12), Inches(0.42), "Do this once. Then skip it next time.", 22, True, INK)
    txt(s, Inches(0.6), Inches(1.52), Inches(12.1), Inches(0.4),
        "1  Install Cursor from cursor.com and sign in.    2  Open chat (Cmd+I) and set mode to Agent.    3  Turn on the Chrome tick below.",
        13, False, MUTED)
    round_box(s, Inches(0.6), Inches(2.05), Inches(6.05), Inches(4.7), WHITE, LINE)
    txt(s, Inches(0.8), Inches(2.18), Inches(5.6), Inches(0.3), "Cursor  ·  Agent, not Ask", 13, True, BLUE)
    pic(s, "shot_cursor_agent.png", Inches(0.8), Inches(2.52), Inches(5.65), max_h=Inches(4.0))
    round_box(s, Inches(6.85), Inches(2.05), Inches(5.9), Inches(4.7), WHITE, LINE)
    txt(s, Inches(7.05), Inches(2.18), Inches(5.5), Inches(0.3), "Chrome  ·  View → Developer", 13, True, BLUE)
    pic(s, "shot_chrome_menu.png", Inches(7.05), Inches(2.52), Inches(5.5), max_h=Inches(4.0))
    footer(s, 4)

    # 5 How-to with screenshots
    s = prs.slides.add_slide(blank)
    light_bg(s)
    logos(s)
    txt(s, Inches(0.6), Inches(0.78), Inches(12), Inches(0.24), "EVERY LIST AFTER THAT", 12, True, BLUE)
    txt(s, Inches(0.6), Inches(1.04), Inches(12), Inches(0.4), "Three steps. Follow the pictures.", 22, True, INK)

    caps = [
        ("1  Open the list", "shot_sn_list.png", "Lead lists → your list. Check the total."),
        ("2  Ask Cursor", "shot_cursor_agent.png", "Agent mode. Paste the sentence. Return."),
        ("3  Check Downloads", "shot_downloads.png", "Open the CSV. Rows must match the total."),
    ]
    for i, (title, img, body) in enumerate(caps):
        x = Inches(0.55 + i * 4.22)
        round_box(s, x, Inches(1.52), Inches(4.05), Inches(3.55), WHITE, LINE)
        txt(s, x + Inches(0.18), Inches(1.6), Inches(3.7), Inches(0.32), title, 14, True, INK)
        pic(s, img, x + Inches(0.18), Inches(1.96), Inches(3.7), max_h=Inches(2.45))
        txt(s, x + Inches(0.18), Inches(4.55), Inches(3.7), Inches(0.4), body, 12, False, MUTED)
        if i < 2:
            chevron(s, x + Inches(3.92), Inches(3.1))

    round_box(s, Inches(0.55), Inches(5.22), Inches(12.2), Inches(1.55), NAVY)
    txt(s, Inches(0.8), Inches(5.36), Inches(11.7), Inches(0.26), "Paste this in Cursor", 12, True, RGBColor(0x7E, 0xC8, 0xFF))
    txt(s, Inches(0.8), Inches(5.68), Inches(11.7), Inches(0.42),
        "Export the open Sales Navigator lead list to CSV. Save it in Downloads.", 18, True, WHITE)
    txt(s, Inches(0.8), Inches(6.18), Inches(11.7), Inches(0.35),
        "If Cursor says the Chrome switch is off, do setup step 3, then reply: go",
        13, False, RGBColor(0xC5, 0xD4, 0xE4))
    footer(s, 5)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
