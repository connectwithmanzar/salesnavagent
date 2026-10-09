#!/usr/bin/env python3
"""Download a Sales Navigator people list from the open Chrome session."""

from __future__ import annotations

import csv
import json
import re
import subprocess
import time
from pathlib import Path
from urllib.parse import parse_qs, urlencode, urlparse, urlunparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path("/Users/manzar.i/Desktop/sdr-data-consolidator")
TMP = ROOT / "tmp_export"
LIST_AS = TMP / "list_sales_tabs.applescript"
JS_PATH = TMP / "extract_list_page.js"
LIST_ID = "7514252219261751296"
BASE = f"https://www.linkedin.com/sales/lists/people/{LIST_ID}?sortCriteria=LAST_ACTIVITY&sortOrder=DESCENDING"
OUT_XLSX = ROOT / "outputs/final/SN_people_list.xlsx"
OUT_CSV = ROOT / "outputs/final/SN_people_list.csv"
DL_XLSX = Path("/Users/manzar.i/Downloads/SN_people_list.xlsx")
DL_CSV = Path("/Users/manzar.i/Downloads/SN_people_list.csv")


def osa(path: Path) -> str:
    r = subprocess.run(["osascript", str(path)], capture_output=True, text=True)
    if r.returncode:
        raise RuntimeError((r.stderr or r.stdout).strip())
    return r.stdout.strip()


def find_tab() -> tuple[int, int] | None:
    raw = osa(LIST_AS)
    for line in raw.splitlines():
        if LIST_ID in line:
            parts = line.split(",", 3)
            return int(parts[0]), int(parts[1])
    for line in raw.splitlines():
        if "linkedin.com/sales" in line:
            parts = line.split(",", 3)
            return int(parts[0]), int(parts[1])
    return None


def open_or_nav(url: str) -> tuple[int, int]:
    found = find_tab()
    if found:
        win, tab = found
    else:
        p = TMP / "open_sn_list.applescript"
        p.write_text(
            '''tell application "Google Chrome"
  tell window 1
    make new tab with properties {URL:"https://www.linkedin.com/sales/home"}
    set idx to count of tabs
  end tell
  return "1," & idx
end tell
'''
        )
        raw = osa(p).strip()
        win, tab = [int(x) for x in raw.split(",")]
        time.sleep(2)
    quoted = url.replace('"', '\\"')
    p = TMP / "nav_sn_list.applescript"
    p.write_text(
        f'tell application "Google Chrome"\n set URL of tab {tab} of window {win} to "{quoted}"\n return "OK"\nend tell\n'
    )
    osa(p)
    return win, tab


def js(win: int, tab: int) -> dict:
    p = TMP / "extract_sn_list.applescript"
    p.write_text(
        f'''tell application "Google Chrome"
  set t to tab {tab} of window {win}
  set js to (do shell script "cat " & quoted form of "{JS_PATH}")
  return execute t javascript js
end tell
'''
    )
    raw = osa(p).strip()
    if not raw or raw == "missing value":
        return {}
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        return {"_raw": raw[:500]}


def page_url(n: int) -> str:
    parsed = urlparse(BASE)
    q = parse_qs(parsed.query)
    q["page"] = [str(n)]
    return urlunparse(parsed._replace(query=urlencode(q, doseq=True)))


def scrape_page(win: int, tab: int, n: int) -> dict:
    quoted = page_url(n).replace('"', '\\"')
    p = TMP / "nav_sn_list.applescript"
    p.write_text(
        f'tell application "Google Chrome"\n set URL of tab {tab} of window {win} to "{quoted}"\n return "OK"\nend tell\n'
    )
    osa(p)
    last = {}
    for attempt in range(8):
        time.sleep(3.2 if attempt == 0 else 1.4)
        data = js(win, tab)
        if data:
            last = data
        if data.get("blocked"):
            return data
        if data.get("leads"):
            return data
        if data.get("listName") and attempt >= 4:
            return data
    return last


def split_name(full: str) -> tuple[str, str]:
    parts = [p for p in re.split(r"\s+", (full or "").strip()) if p]
    if not parts:
        return "", ""
    if len(parts) == 1:
        return parts[0], ""
    return parts[0], " ".join(parts[1:])


def write_files(leads: list[dict], list_name: str, total_hint: int) -> None:
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    cols = ["First Name", "Last Name", "Full Name", "Title", "Company", "Location", "Sales Navigator URL", "Date added"]
    rows = []
    for lead in leads:
        first, last = split_name(lead.get("name") or "")
        rows.append(
            {
                "First Name": first,
                "Last Name": last,
                "Full Name": lead.get("name") or "",
                "Title": lead.get("title") or "",
                "Company": lead.get("company") or "",
                "Location": lead.get("location") or "",
                "Sales Navigator URL": lead.get("linkedin_url") or "",
                "Date added": lead.get("extra") or "",
            }
        )

    with OUT_CSV.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    navy = PatternFill("solid", fgColor="1B365D")
    white = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    body = Font(name="Arial", size=10)
    wrap = Alignment(vertical="center", wrap_text=True)
    thin = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    alt = PatternFill("solid", fgColor="F8FAFC")
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(1, c)
        cell.fill = navy
        cell.font = white
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append([r[c] for c in cols])
        rr = ws.max_row
        for c in range(1, len(cols) + 1):
            cell = ws.cell(rr, c)
            cell.font = body
            cell.alignment = wrap
            cell.border = thin
            if rr % 2 == 0:
                cell.fill = alt
        ws.row_dimensions[rr].height = 22
    last = max(ws.max_row, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    for col, w in zip("ABCDEFGH", [16, 18, 26, 42, 28, 22, 56, 14]):
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    wb.properties.title = list_name or "SN people list"
    wb.save(OUT_XLSX)
    try:
        DL_XLSX.write_bytes(OUT_XLSX.read_bytes())
        DL_CSV.write_bytes(OUT_CSV.read_bytes())
    except OSError as exc:
        print("downloads copy failed", exc)
    print(f"wrote {len(rows)} leads ({list_name or 'list'}; SN said {total_hint})", flush=True)
    print(str(DL_XLSX), flush=True)


def main() -> None:
    win, tab = open_or_nav(BASE)
    print(f"SN window={win} tab={tab}", flush=True)
    all_leads = []
    seen = set()
    list_name = ""
    total_hint = 0
    page = 1
    empty = 0
    while page <= 40:
        print(f"page {page}...", flush=True)
        data = scrape_page(win, tab, page)
        if data.get("blocked"):
            raise SystemExit("LinkedIn login wall — keep Sales Navigator open in Chrome and retry")
        if data.get("listName"):
            list_name = data["listName"]
        if data.get("total"):
            total_hint = int(data["total"])
        leads = data.get("leads") or []
        print(f"  {data.get('count', 0)} rows | list={list_name!r} total={total_hint} href={str(data.get('href') or '')[:80]}", flush=True)
        if not leads:
            empty += 1
            if empty >= 2:
                break
            page += 1
            continue
        empty = 0
        new = 0
        for lead in leads:
            key = (lead.get("linkedin_url") or lead.get("name") or "").lower()
            if not key or key in seen:
                continue
            seen.add(key)
            all_leads.append(lead)
            new += 1
        print(f"  +{new} unique (running {len(all_leads)})", flush=True)
        if total_hint and len(all_leads) >= total_hint:
            break
        if new == 0:
            break
        page += 1
    if not all_leads:
        print("sample", json.dumps(data.get("sample") or data, indent=2)[:1500], flush=True)
        raise SystemExit("No leads extracted")
    write_files(all_leads, list_name, total_hint)


if __name__ == "__main__":
    main()
